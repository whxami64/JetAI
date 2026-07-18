"""Tests for the xlsx/docx/pdf preprocessing pipeline."""

from __future__ import annotations

import shutil
from pathlib import Path

import docx
import openpyxl
import pandas as pd
import pytest
import xlwt

from jetai.dataset import build_inventory
from jetai.preprocess import preprocess_dataset

_SAMPLE_PDF = (
    Path(__file__).parent.parent
    / "data"
    / "Uebungsdaten Muster Verpackungen"
    / "Begleitdokumente"
    / "JA-Entwurf_2025_Auszug_Bilanz_GuV.pdf"
)


@pytest.fixture
def sample_dataset(tmp_path: Path) -> Path:
    root = tmp_path / "Uebungsdaten Muster"
    (root / "Begleitdokumente").mkdir(parents=True)

    single_sheet = openpyxl.Workbook()
    single_sheet.active.append(["Konto", "Saldo"])
    single_sheet.active.append(["100000", "1.234,56"])
    single_sheet.save(root / "Begleitdokumente" / "Saldenliste_2025.xlsx")

    multi_sheet = openpyxl.Workbook()
    multi_sheet.active.title = "Sheet A"
    multi_sheet.active.append(["x"])
    multi_sheet.create_sheet("Sheet B").append(["y"])
    multi_sheet.save(root / "Begleitdokumente" / "OP-Liste_2025.xlsx")

    legacy_workbook = xlwt.Workbook()
    legacy_sheet = legacy_workbook.add_sheet("Sheet1")
    legacy_sheet.write(0, 0, "Konto")
    legacy_sheet.write(0, 1, "Saldo")
    legacy_sheet.write(1, 0, "200000")
    legacy_sheet.write(1, 1, "9.876,50")
    legacy_workbook.save(str(root / "Begleitdokumente" / "Saldenliste_2024.xls"))

    paper = docx.Document()
    paper.add_heading("Pruefungsplanung", level=1)
    paper.add_paragraph("Wesentlichkeit 400.000 EUR.")
    paper.save(root / "Begleitdokumente" / "Pruefungsplanung_JET_2025.docx")

    if _SAMPLE_PDF.exists():
        shutil.copy(_SAMPLE_PDF, root / "Begleitdokumente" / "JA-Entwurf_2025.pdf")

    return root


def test_preprocess_converts_xlsx_to_csv(sample_dataset: Path) -> None:
    result = preprocess_dataset(sample_dataset)

    single = sample_dataset / "Begleitdokumente" / "Saldenliste_2025.csv"
    assert single in result.converted
    frame = pd.read_csv(single, sep=";")
    assert list(frame.columns) == ["Konto", "Saldo"]


def test_preprocess_splits_multi_sheet_xlsx(sample_dataset: Path) -> None:
    result = preprocess_dataset(sample_dataset)

    names = {p.name for p in result.converted}
    assert "OP-Liste_2025__Sheet_A.csv" in names
    assert "OP-Liste_2025__Sheet_B.csv" in names


def test_preprocess_converts_legacy_xls_to_csv(sample_dataset: Path) -> None:
    result = preprocess_dataset(sample_dataset)

    out = sample_dataset / "Begleitdokumente" / "Saldenliste_2024.csv"
    assert out in result.converted
    frame = pd.read_csv(out, sep=";")
    assert list(frame.columns) == ["Konto", "Saldo"]
    assert not (sample_dataset / "Begleitdokumente" / "Saldenliste_2024.xls").exists()


def test_preprocess_converts_docx_to_markdown(sample_dataset: Path) -> None:
    preprocess_dataset(sample_dataset)

    out = sample_dataset / "Begleitdokumente" / "Pruefungsplanung_JET_2025.md"
    assert out.exists()
    assert "Wesentlichkeit 400.000 EUR" in out.read_text(encoding="utf-8")


@pytest.mark.skipif(not _SAMPLE_PDF.exists(), reason="sample pdf not present in repo")
def test_preprocess_keeps_pdf_labels_and_values_on_one_line(sample_dataset: Path) -> None:
    preprocess_dataset(sample_dataset)

    out = sample_dataset / "Begleitdokumente" / "JA-Entwurf_2025.md"
    text = out.read_text(encoding="utf-8")
    assert "Sachanlagen 19.729.014,76" in text


def test_preprocess_archives_originals_to_stale(sample_dataset: Path) -> None:
    result = preprocess_dataset(sample_dataset)

    assert not (sample_dataset / "Begleitdokumente" / "Saldenliste_2025.xlsx").exists()
    archived = sample_dataset / "stale" / "Begleitdokumente" / "Saldenliste_2025.xlsx"
    assert archived in result.archived
    assert archived.exists()


def test_preprocess_is_idempotent(sample_dataset: Path) -> None:
    first = preprocess_dataset(sample_dataset)
    second = preprocess_dataset(sample_dataset)

    assert first.converted
    assert not second.converted
    assert not second.archived


def test_stale_directory_excluded_from_inventory(sample_dataset: Path) -> None:
    preprocess_dataset(sample_dataset)
    inventory = build_inventory(sample_dataset)

    assert not any(p.suffix.lower() in {".xlsx", ".xls"} for p in inventory.tables)
    assert not any(p.suffix.lower() in {".docx", ".pdf"} for p in inventory.documents)
